"""DOCX and JSON exports (spec §10, §14)."""

from __future__ import annotations

import io
import json

from docx import Document
from sqlalchemy.orm import Session

from export.docx import to_docx_bytes
from export.json_export import to_dict, to_json_bytes
from models.enums import Destination
from tests._export_helpers import profile_of, seeded_cv


def test_docx_structure_without_tables(session: Session) -> None:
    cv = seeded_cv(session)
    data = to_docx_bytes(cv, profile_of(cv, Destination.ATS))
    doc = Document(io.BytesIO(data))
    assert doc.tables == []
    styles = [p.style.name for p in doc.paragraphs if p.style is not None]
    assert styles[0] == "Title"
    assert {"Heading 1", "Heading 2", "List Bullet", "Normal"} <= set(styles)
    headings = [
        p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"
    ]
    assert headings == [
        "RESUMEN PROFESIONAL", "EXPERIENCIA PROFESIONAL", "EDUCACIÓN", "HABILIDADES",
        "PROYECTOS", "CERTIFICACIONES", "IDIOMAS",
    ]  # fmt: skip
    props = doc.core_properties
    assert props.author == "Andrés Felipe Rojas Martínez"
    assert props.subject == "Hoja de vida"
    assert props.title.endswith("Hoja de vida")
    assert len(doc.sections) == 1


def test_docx_contains_full_text(session: Session) -> None:
    cv = seeded_cv(session)
    doc = Document(io.BytesIO(to_docx_bytes(cv, profile_of(cv, Destination.ATS))))
    text = "\n".join(p.text for p in doc.paragraphs)
    expected = [cv.basics.full_name, cv.basics.email, "+57 300 123 4567", cv.basics.summary or ""]
    for exp in cv.experiences:
        expected += [exp.company, exp.role, *(a.text for a in exp.achievements)]
    expected += [e.institution for e in cv.education]
    expected += [s.name for c in cv.skill_categories for s in c.skills]
    expected += [p.name for p in cv.projects] + [c.name for c in cv.certifications]
    expected += [lang.name for lang in cv.languages]
    missing = [e for e in expected if e not in text]
    assert missing == []


def test_docx_mail_profile_uses_title_case_headings(session: Session) -> None:
    cv = seeded_cv(session)
    doc = Document(io.BytesIO(to_docx_bytes(cv, profile_of(cv, Destination.MAIL))))
    headings = [
        p.text for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"
    ]
    assert "Experiencia profesional" in headings
    assert doc.tables == []


def test_json_export_is_resume_out_dump(session: Session) -> None:
    cv = seeded_cv(session)
    data = json.loads(to_json_bytes(cv).decode("utf-8"))
    assert data == to_dict(cv) == cv.model_dump(mode="json")
    assert "Andrés" in to_json_bytes(cv).decode("utf-8")  # accents kept
