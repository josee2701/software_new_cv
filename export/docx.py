"""DOCX export (spec §10 step 7): same structure as the ``ats`` template.

One column, no tables, built-in styles only: ``Title`` (name), ``Heading 1`` (sections),
``Heading 2`` (items), ``List Bullet`` (achievements) and ``Normal``. Content comes from the same
render context as the HTML (``render.context.build_context``), so texts and order match.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.shared import Mm, Pt

from render.context import build_context
from schemas import ExportProfileOut, ResumeOut

FONT_SIZES = {"compact": 9.5, "normal": 10.5, "wide": 11.5}
MARGINS_MM = {"compact": 14, "normal": 18, "wide": 22}
SUBJECT = "Hoja de vida"


def _setup(doc: DocxDocument, ctx: dict[str, Any]) -> None:
    density = ctx["profile"]["density"]
    font = ctx["profile"]["font_family"] if ctx["destination"] == "ats" else "Arial"
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = doc.styles[style_name]
        style.font.name = font
    doc.styles["Normal"].font.size = Pt(FONT_SIZES.get(density, 10.5))
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(215.9), Mm(279.4)  # Letter
    margin = Mm(MARGINS_MM.get(density, 18))
    section.left_margin = section.right_margin = margin
    section.top_margin = section.bottom_margin = margin
    props = doc.core_properties
    props.title = ctx["doc_title"]
    props.author = ctx["basics"]["full_name"]
    props.subject = SUBJECT
    props.comments = "cv-ats"
    props.last_modified_by = "cv-ats"


def _labeled(doc: DocxDocument, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(text)


def _meta(doc: DocxDocument, *parts: str) -> None:
    text = " · ".join(p for p in parts if p)
    if text:
        doc.add_paragraph(text)


def _technologies(doc: DocxDocument, ctx: dict[str, Any], items: list[str]) -> None:
    if items:
        _labeled(doc, ctx["t"]["technologies"], ", ".join(items))


def _section(doc: DocxDocument, ctx: dict[str, Any], section: dict[str, Any]) -> None:
    key, t = section["key"], ctx["t"]
    doc.add_heading(section["title"], level=1)
    if key == "summary":
        doc.add_paragraph(ctx["basics"]["summary"])
        return
    for item in section["entries"]:
        if key == "experience":
            doc.add_heading(f"{item['role']} — {item['company']}", level=2)
            _meta(doc, item["location"], item["dates"])
            if item["description"]:
                doc.add_paragraph(item["description"])
            for achievement in item["achievements"]:
                doc.add_paragraph(achievement, style="List Bullet")
            _technologies(doc, ctx, item["technologies"])
        elif key == "education":
            doc.add_heading(f"{item['degree']} — {item['institution']}", level=2)
            _meta(doc, item["location"], item["dates"])
            gpa = f"{t['gpa']}: {item['gpa']}" if item["gpa"] else ""
            _meta(doc, item["field"], gpa)
            if item["description"]:
                doc.add_paragraph(item["description"])
        elif key == "skills":
            _labeled(doc, item["category"], item["names"])
        elif key == "projects":
            title = item["name"] + (f" — {item['role']}" if item["role"] else "")
            doc.add_heading(title, level=2)
            _meta(doc, item["dates"], item["url_display"])
            doc.add_paragraph(item["description"])
            _technologies(doc, ctx, item["technologies"])
        elif key == "certifications":
            doc.add_heading(f"{item['name']} — {item['issuer']}", level=2)
            expiry = f"{t['expires']}: {item['expiry']}" if item["expiry"] else ""
            credential = (
                f"{t['credential']}: {item['credential_id']}" if item["credential_id"] else ""
            )
            _meta(doc, item["date"], expiry, credential, item["url_display"])
        elif key == "languages":
            _labeled(doc, item["name"], item["level_label"])


def build_document(cv: ResumeOut, profile: ExportProfileOut) -> DocxDocument:
    ctx = build_context(cv, profile)
    doc = Document()
    _setup(doc, ctx)
    basics = ctx["basics"]
    doc.add_paragraph(basics["full_name"], style="Title")
    if basics["headline"]:
        doc.add_paragraph(basics["headline"])
    contact = [basics["email"], basics["phone_display"], basics["location"]]
    contact += [link["display"] for link in ctx["links"]]
    doc.add_paragraph(" | ".join(c for c in contact if c))
    for section in ctx["sections"]:
        _section(doc, ctx, section)
    return doc


def to_docx_bytes(cv: ResumeOut, profile: ExportProfileOut) -> bytes:
    buffer = io.BytesIO()
    build_document(cv, profile).save(buffer)
    return buffer.getvalue()
