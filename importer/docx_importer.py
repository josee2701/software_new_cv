"""DOCX → lines (paragraphs with style, bold and list information)."""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

from docx import Document
from docx.text.paragraph import Paragraph

from importer.lines import Line
from importer.text import is_bullet_text

_STYLE_SIZES = {"Title": 26.0, "Heading 1": 14.0, "Heading 2": 13.0, "Heading 3": 12.0}


def _size(paragraph: Paragraph) -> float:
    for run in paragraph.runs:
        if run.font.size is not None:
            return float(run.font.size.pt)
    style: Any = paragraph.style
    while style is not None:
        if style.font.size is not None:
            return float(style.font.size.pt)
        style = style.base_style
    name = paragraph.style.name if paragraph.style is not None else ""
    return _STYLE_SIZES.get(name or "", 11.0)


def _is_list(paragraph: Paragraph, style: str) -> bool:
    if style.startswith("List"):
        return True
    props = paragraph._p.pPr
    return props is not None and props.numPr is not None


def _paragraphs(document: Any) -> Iterator[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:  # tables are not produced by cv-ats, but others use them
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def docx_lines(data: bytes) -> list[Line]:
    document = Document(io.BytesIO(data))
    lines: list[Line] = []
    for paragraph in _paragraphs(document):
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        runs = [r for r in paragraph.runs if r.text.strip()]
        bold = style == "Title" or style.startswith("Heading")
        bold = bold or (bool(runs) and all(bool(r.bold) for r in runs))
        bullet = _is_list(paragraph, style)
        if bullet and not is_bullet_text(text):
            text = f"• {text}"
        indent = paragraph.paragraph_format.left_indent
        lines.append(
            Line(
                text=text,
                size=_size(paragraph),
                bold=bold,
                x0=float(indent.pt) if indent is not None else 0.0,
                style=style,
                bullet=bullet,
            )
        )
    return lines
